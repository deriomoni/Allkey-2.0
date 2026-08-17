"""Personnel module API — STATELESS document generation.

The server never stores employees' personal data: prikaz / preview / zayavlenie
all render from the data POSTed in the request body (the client keeps it in the
form draft). Only rates and ИИН/БИН validation are read-only helpers. Everything
is gated by the `hr` service (require_service), the module's security boundary.

PII rule (ТЗ §9): ИИН and document numbers arrive only in the request body, never
in path/query params, and are not logged.
"""
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse

from app.personnel import inventory_columns as ic

from app.personnel.context import (
    build_order_context, build_deduction_application_context, build_prikaz_preview,
    build_matotvet_context, build_nekonkurencii_context, build_akt_context,
    build_perechen_context, build_trudovoy_context,
    build_polozhenie_pd_context, build_prikaz_pd_context, DEDUCTION_TEXTS,
)
from app.personnel.generator import render_template, render_bytes, build_zip, output_filename
from app.personnel.helpers.iin import is_valid_iin, is_valid_bin, parse_iin
from app.personnel.helpers.validation import (
    validate_iin_matches, validate_salary, validate_probation, validate_dates,
)
from app.personnel.rates import get_rates
from app.personnel.schemas import (
    IinCheckRequest, IinCheckResponse,
    BinCheckRequest, BinCheckResponse,
    RatesResponse, PrikazRequest, ZayavlenieVychetyRequest, PrikazPreviewResponse,
    PackageRequest, InventoryParseResponse, InventoryItemIn, InventoryColumn,
    TrudovoyRequest,
)
from app.services.dependencies import require_service
from app.users.models import User

SERVICE_CODE = "hr"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP_MEDIA_TYPE = "application/zip"

router = APIRouter(prefix="/personnel", tags=["personnel"])


def _docx_response(document, filename: str) -> StreamingResponse:
    # RFC 5987 for the non-ASCII (Cyrillic) filename.
    disposition = f"attachment; filename*=UTF-8''{quote(filename)}"
    return StreamingResponse(
        document, media_type=DOCX_MEDIA_TYPE,
        headers={"Content-Disposition": disposition},
    )


def _hiring_warnings(employee, employment) -> List[str]:
    """Soft warnings for the form (salary rate-aware, probation, dates, ИИН match)."""
    warnings: List[str] = []
    if employment.salary is not None:
        rate = float(employment.rate) if employment.rate is not None else 1.0
        warnings += validate_salary(employment.salary, rate)
    warnings += validate_probation(employment.probation_months or 0)
    if employment.start_date and employment.contract_date:
        warnings += validate_dates(employment.start_date, employment.contract_date)
    if employee.iin:
        warnings += validate_iin_matches(employee.iin, employee.birth_date, employee.gender)
    return warnings


@router.get("/rates", response_model=RatesResponse)
async def current_rates(
    on: date | None = None,
    _user: User = Depends(require_service(SERVICE_CODE)),
):
    """Payroll rates in force on `on` (default: today). Feeds the calculator."""
    r = get_rates(on)
    return RatesResponse(
        effective_from=r.effective_from, effective_to=r.effective_to,
        mrp=r.mrp, mzp=r.mzp, ipn_rate=r.ipn_rate, base_deduction_mrp=r.base_deduction_mrp,
        opv_rate=r.opv_rate, opvr_rate=r.opvr_rate, so_rate=r.so_rate,
        vosms_rate=r.vosms_rate, oosms_rate=r.oosms_rate, sn_rate=r.sn_rate,
        unified_payment_rate=r.unified_payment_rate,
    )


@router.post("/validate/iin", response_model=IinCheckResponse)
async def validate_iin(
    data: IinCheckRequest,
    _user: User = Depends(require_service(SERVICE_CODE)),
):
    """Validate an ИИН and decode birth date / gender, warning on mismatches."""
    valid = is_valid_iin(data.iin)
    info = parse_iin(data.iin)
    warnings = validate_iin_matches(data.iin, data.birth_date, data.gender) if valid else []
    return IinCheckResponse(
        valid=valid,
        birth_date=info.birth_date if info else None,
        gender=info.gender if info else None,
        warnings=warnings,
    )


@router.post("/validate/bin", response_model=BinCheckResponse)
async def validate_bin(
    data: BinCheckRequest,
    _user: User = Depends(require_service(SERVICE_CODE)),
):
    """Validate a БИН (length + checksum)."""
    return BinCheckResponse(valid=is_valid_bin(data.bin))


@router.post("/documents/prikaz/preview", response_model=PrikazPreviewResponse)
async def prikaz_preview(
    data: PrikazRequest,
    user: User = Depends(require_service(SERVICE_CODE)),
):
    """Editable-fields preview for the form (stateless): full rendered context, the
    editable auto-values (ФИО падежи + должность + сумма прописью) and soft warnings."""
    preview = build_prikaz_preview(
        data.company, data.employee, data.employment,
        hr_responsible_fio=data.hr_responsible_fio or user.full_name,
    )
    preview["warnings"] = _hiring_warnings(data.employee, data.employment)
    return PrikazPreviewResponse(**preview)


@router.post("/documents/prikaz")
async def generate_prikaz(
    data: PrikazRequest,
    user: User = Depends(require_service(SERVICE_CODE)),
):
    """Render «Приказ о приёме на работу» .docx from the posted data (§4.4).

    Fully stateless: nothing is read from or written to the database."""
    context = build_order_context(
        data.company, data.employee, data.employment,
        hr_responsible_fio=data.hr_responsible_fio or user.full_name,
    )
    document = render_template("prikaz_o_prieme.docx", context)
    filename = output_filename(
        data.employee.last_name, data.employee.first_name, data.employee.middle_name or "",
        "ПриказПриём", data.employment.order_date,
    )
    return _docx_response(document, filename)


@router.post("/documents/zayavlenie-vychety")
async def generate_zayavlenie_vychety(
    data: ZayavlenieVychetyRequest,
    _user: User = Depends(require_service(SERVICE_CODE)),
):
    """Render «Заявление о применении налоговых вычетов» .docx from the posted data (§4.5)."""
    unknown = [k for k in data.deductions if k not in DEDUCTION_TEXTS]
    if unknown:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Неизвестные виды вычета: {', '.join(unknown)}",
        )

    apply_from = data.apply_from or data.employment.start_date
    context = build_deduction_application_context(
        data.company, data.employee, data.deductions,
        apply_from=apply_from,
        application_date=data.employment.application_date or apply_from,
    )
    document = render_template("zayavlenie_vychety_ipn.docx", context)
    filename = output_filename(
        data.employee.last_name, data.employee.first_name, data.employee.middle_name or "",
        "ЗаявлениеВычеты", data.employment.application_date,
    )
    return _docx_response(document, filename)


@router.post("/documents/trudovoy")
async def generate_trudovoy(
    data: TrudovoyRequest,
    _user: User = Depends(require_service(SERVICE_CODE)),
):
    """Render the bilingual «Трудовой договор» (+ Приложение №1 об окладе) .docx
    from the posted data (§4.2). Kazakh number/date/term forms are computed; the
    Kazakh column must be proofread before client use (see README)."""
    context = build_trudovoy_context(data.company, data.employee, data.employment, data.contract)
    document = render_template("trudovoy_dogovor.docx", context)
    filename = output_filename(
        data.employee.last_name, data.employee.first_name, data.employee.middle_name or "",
        "ТрудовойДоговор", data.contract.doc_date,
    )
    return _docx_response(document, filename)


@router.post("/documents/package")
async def generate_package(
    data: PackageRequest,
    user: User = Depends(require_service(SERVICE_CODE)),
):
    """Render the selected documents (ТЗ §2 галочки) into one ZIP, statelessly.

    Each document key requires its own inputs; a selected document missing its
    inputs is a 400 rather than a silently empty file."""
    if not data.documents:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Не выбран ни один документ")

    emp = data.employee

    def fname(label: str, on) -> str:
        return output_filename(emp.last_name, emp.first_name, emp.middle_name or "", label, on)

    def need(value, detail: str):
        if value is None or value == []:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=detail)
        return value

    files: List = []
    for doc in data.documents:
        if doc == "prikaz":
            ctx = build_order_context(data.company, emp, data.employment,
                                      hr_responsible_fio=data.hr_responsible_fio or user.full_name)
            files.append((fname("ПриказПриём", data.employment.order_date),
                          render_bytes("prikaz_o_prieme.docx", ctx)))
        elif doc == "zayavlenie":
            unknown = [k for k in data.deductions if k not in DEDUCTION_TEXTS]
            if unknown:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                                    detail=f"Неизвестные виды вычета: {', '.join(unknown)}")
            apply_from = data.apply_from or data.employment.start_date
            ctx = build_deduction_application_context(
                data.company, emp, data.deductions,
                apply_from=apply_from,
                application_date=data.employment.application_date or apply_from,
            )
            files.append((fname("ЗаявлениеВычеты", data.employment.application_date),
                          render_bytes("zayavlenie_vychety_ipn.docx", ctx)))
        elif doc == "matotvet":
            need(data.liability, "Для договора о матответственности нужны реквизиты (liability)")
            ctx = build_matotvet_context(data.company, emp, data.employment, data.liability)
            files.append((fname("ДоговорМатОтв", data.liability.doc_date),
                          render_bytes("dogovor_matotvetstvennost.docx", ctx)))
        elif doc == "akt":
            need(data.act, "Для акта приёма-передачи нужны данные акта (act)")
            need(data.inventory, "Для акта нужна опись позиций (inventory)")
            ctx = build_akt_context(data.company, emp, data.employment, data.act, data.inventory, data.liability)
            files.append((fname("АктПриёмаПередачи", data.act.doc_date),
                          render_bytes("akt_priema_peredachi.docx", ctx)))
        elif doc == "nekonkurencii":
            need(data.noncompete, "Для договора о неконкуренции нужны условия (noncompete)")
            ctx = build_nekonkurencii_context(data.company, emp, data.employment, data.noncompete)
            files.append((fname("ДоговорНеконкуренции", data.noncompete.doc_date),
                          render_bytes("dogovor_nekonkurencii.docx", ctx)))
        elif doc == "perechen":
            need(data.perechen, "Для приказа об утверждении перечня нужны данные (perechen)")
            ctx = build_perechen_context(data.company, data.perechen, data.noncompete)
            files.append((fname("ПриказПеречень", data.perechen.doc_date),
                          render_bytes("prikaz_perechen_nekonkurencii.docx", ctx)))
        elif doc == "td":
            need(data.contract, "Для трудового договора нужны данные договора (contract)")
            ctx = build_trudovoy_context(data.company, data.employee, data.employment, data.contract)
            files.append((fname("ТрудовойДоговор", data.contract.doc_date),
                          render_bytes("trudovoy_dogovor.docx", ctx)))
        elif doc == "polozhenie_pd":
            need(data.policy, "Для Положения о ПД нужны данные приказа (policy)")
            ctx = build_polozhenie_pd_context(data.company, data.policy)
            files.append((fname("ПоложениеПД", data.policy.doc_date),
                          render_bytes("polozhenie_personalnye_dannye.docx", ctx)))
        elif doc == "prikaz_pd":
            need(data.policy, "Для приказа об ответственном за ПД нужны данные (policy)")
            ctx = build_prikaz_pd_context(data.company, data.policy)
            files.append((fname("ПриказОтветственныйПД", data.policy.doc_date),
                          render_bytes("prikaz_otvetstvennyy_pd.docx", ctx)))
        else:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Неизвестный документ: {doc}")

    archive = build_zip(files)
    zip_name = fname("Пакет", data.employment.order_date).replace(".docx", ".zip")
    disposition = f"attachment; filename*=UTF-8''{quote(zip_name)}"
    return StreamingResponse(archive, media_type=ZIP_MEDIA_TYPE, headers={"Content-Disposition": disposition})


def _to_decimal(value) -> Decimal:
    if value is None:
        return Decimal(0)
    text = str(value).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal(0)


@router.post("/parse-inventory", response_model=InventoryParseResponse)
async def parse_inventory(
    file: UploadFile = File(...),
    header_row: Optional[int] = Form(None),
    col_name: Optional[int] = Form(None),
    col_qty: Optional[int] = Form(None),
    col_price: Optional[int] = Form(None),
    col_code: Optional[int] = Form(None),
    col_unit: Optional[int] = Form(None),
    _user: User = Depends(require_service(SERVICE_CODE)),
):
    """Parse an .xlsx опись into inventory rows (name/code/unit/qty/price).

    Column titles are matched against synonyms (inventory_columns.COLUMN_SYNONYMS)
    after normalization, and the header row is searched below any 1С preamble
    (org name / period / blank rows). If auto-detection fails, returns
    status='needs_mapping' with the file's columns so the client shows a mapping
    screen; a repeat call with col_name/col_qty/col_price (column indices) parses
    explicitly. Stateless — the file is parsed in memory and discarded."""
    from openpyxl import load_workbook

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="Не удалось прочитать файл (нужен .xlsx)")
    rows = [r for r in wb.active.iter_rows(values_only=True)]
    if not rows:
        return InventoryParseResponse(status="parsed", items=[])

    def cell(row, i):
        return row[i] if (i is not None and i < len(row) and row[i] is not None) else ""

    def parse(header_idx: int, mapping: dict) -> List[InventoryItemIn]:
        out: List[InventoryItemIn] = []
        for row in rows[header_idx + 1:]:
            name = str(cell(row, mapping.get("name"))).strip()
            if not name:
                continue
            out.append(InventoryItemIn(
                name=name,
                code=str(cell(row, mapping.get("code"))).strip(),
                unit=str(cell(row, mapping.get("unit"))).strip(),
                qty=_to_decimal(cell(row, mapping.get("qty"))),
                price=_to_decimal(cell(row, mapping.get("price"))),
            ))
        return out

    # 1) Explicit mapping supplied by the manual-mapping step.
    if col_name is not None and col_qty is not None and col_price is not None:
        mapping = {"name": col_name, "qty": col_qty, "price": col_price}
        if col_code is not None:
            mapping["code"] = col_code
        if col_unit is not None:
            mapping["unit"] = col_unit
        idx = header_row if header_row is not None else 0
        return InventoryParseResponse(status="parsed", header_row=idx, items=parse(idx, mapping))

    # 2) Auto-detect header + columns via synonyms.
    idx, mapping = ic.find_header(rows)
    if mapping is not None:
        return InventoryParseResponse(status="parsed", header_row=idx, items=parse(idx, mapping))

    # 3) Could not recognize — offer manual mapping instead of rejecting.
    hr = ic.guess_header_row(rows)
    header = list(rows[hr]) if hr < len(rows) else []
    data = [r for r in rows[hr + 1:] if any(str(c).strip() for c in r if c is not None)]
    columns = []
    for ci, title in enumerate(header):
        samples = [s for r in data[:3] if (s := str(cell(r, ci)).strip())]
        text = str(title).strip() if title is not None and str(title).strip() else f"Колонка {ci + 1}"
        columns.append(InventoryColumn(index=ci, title=text, samples=samples))
    return InventoryParseResponse(status="needs_mapping", header_row=hr, columns=columns)
