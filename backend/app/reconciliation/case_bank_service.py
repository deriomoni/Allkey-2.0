import re
import math
import datetime as dt
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from app.reconciliation.aoa import read_aoa

Cell = Any


def _mathround(x: float) -> int:
    return math.floor(x + 0.5) if x >= 0 else math.ceil(x - 0.5)


def _r2(n: float) -> float:
    return _mathround(n * 100) / 100


def _fmt_date(d: int, m: int, y: int) -> str:
    return f"{d:02d}.{m:02d}.{y}"


def _parse_date(v: Cell) -> Optional[str]:
    if isinstance(v, (dt.datetime, dt.date)):
        return _fmt_date(v.day, v.month, v.year)
    if isinstance(v, str):
        m = re.match(r"^\s*(\d{1,2})[.\/-](\d{1,2})[.\/-](\d{2,4})", v.strip())
        if m:
            y = m.group(3)
            y = f"20{y}" if len(y) == 2 else y
            return _fmt_date(int(m.group(1)), int(m.group(2)), int(y))
    return None


def _date_key(d: str) -> int:
    dd, mm, yy = (int(x) for x in d.split("."))
    return yy * 10000 + mm * 100 + dd


def _parse_num(v: Cell) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(float(v)) else None
    if not isinstance(v, str):
        return None
    s = re.sub(r"[\s\u00a0\u202f]", "", v)
    if not s:
        return None
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    if not re.fullmatch(r"-?\d+(\.\d+)?", s):
        return None
    try:
        n = float(s)
    except ValueError:
        return None
    return n if math.isfinite(n) else None


def _first_line(v: Cell) -> str:
    return v.split("\n")[0].strip() if isinstance(v, str) else ""


def _nth_line(v: Cell, n: int) -> str:
    if not isinstance(v, str):
        return ""
    parts = [x.strip() for x in v.split("\n") if x.strip()]
    return parts[n] if n < len(parts) else ""


def _doc_no_1c(v: Cell) -> str:
    if not isinstance(v, str):
        return ""
    m = re.search(r"№?\s*0*(\d{1,10})\s+от", v)
    return m.group(1) if m else ""


def _find_header_row(aoa: List[List[Cell]], needles: List[str]):
    for i, row in enumerate(aoa):
        if not row:
            continue
        cols: Dict[str, int] = {}
        hit = 0
        for need in needles:
            idx = next((j for j, c in enumerate(row)
                        if isinstance(c, str) and c.strip().lower() == need.lower()), -1)
            if idx >= 0:
                cols[need] = idx
                hit += 1
        if hit >= 2 and "Дебет" in cols and "Кредит" in cols:
            return {"row": i, "cols": cols}
    return None


def _norm(c) -> str:
    return re.sub(r"\s+", " ", str(c)).strip().lower() if isinstance(c, str) else ""


def _amount_col(sub, start, end):
    """Колонка суммы внутри блока «Дебет»/«Кредит».

    В карточке 1С под «Дебет»/«Кредит» идут подстолбцы: «Счёт» (номер счёта),
    иногда «Вал.» (валюта) и собственно сумма. Берём первый столбец блока,
    который НЕ «Счёт» и НЕ валюта. Так корректно работает и тенговая карточка
    (сумма сразу за «Счёт», +1), и валютная (+2, если есть столбец валюты).
    """
    acct = cur = None
    for c in range(start, min(end, len(sub))):
        t = _norm(sub[c])
        if t in ("счет", "счёт"):
            acct = c
        elif t.startswith("вал"):
            cur = c
    for c in range(start, end):
        if c in (acct, cur):
            continue
        return c
    return start + 1


_BANK_ALIASES = {
    "date": ["дата", "күні"],
    "debit": ["дебет"],
    "credit": ["кредит"],
    "cp": ["контрагент", "корреспондент"],
    "pur": ["назначение платежа", "назначение", "мақсаты"],
}


def _find_bank_cols(aoa):
    """Найти строку заголовков выписки и индексы колонок по синонимам (двуязычные шапки)."""
    for i, row in enumerate(aoa):
        if not row:
            continue
        found = {}
        for key, aliases in _BANK_ALIASES.items():
            for j, c in enumerate(row):
                n = _norm(c)
                if n and any(a in n for a in aliases):
                    found[key] = j
                    break
        if "debit" in found and "credit" in found and "date" in found:
            return i, found
    return None, {}


def _signed_balance(row: List[Cell], amount: float) -> float:
    idx = -1
    for j, c in enumerate(row):
        n = _parse_num(c)
        if n is not None and abs(n - amount) < 1e-9:
            idx = j
    if idx < 0:
        return amount
    for j in range(idx - 1, -1, -1):
        c = row[j]
        if isinstance(c, str) and re.fullmatch(r"[ДК]", c.strip()):
            return -amount if (c.strip() == "К" and amount > 0) else amount
    return amount


def _parse_1c(aoa: List[List[Cell]]) -> Dict[str, Any]:
    hdr = _find_header_row(aoa, ["Дебет", "Кредит", "Общий оборот", "Текущее сальдо"])
    if hdr:
        sub = aoa[hdr["row"] + 1] if hdr["row"] + 1 < len(aoa) else []
        deb_h = hdr["cols"]["Дебет"]
        kre_h = hdr["cols"]["Кредит"]
        oborot = hdr["cols"].get("Общий оборот", hdr["cols"].get("Текущее сальдо", kre_h + 3))
        deb_col = _amount_col(sub, deb_h, kre_h)
        kre_col = _amount_col(sub, kre_h, oborot)
    else:
        deb_col, kre_col = 7, 10

    show_col = -1
    if hdr:
        hrow = aoa[hdr["row"]] or []
        show_col = next((j for j, c in enumerate(hrow)
                         if isinstance(c, str) and re.search(r"Показатель", c, re.I)), -1)

    def is_val(row: Optional[List[Cell]]) -> bool:
        return (show_col >= 0 and row is not None and show_col < len(row)
                and isinstance(row[show_col], str) and re.search(r"Вал\.", row[show_col]) is not None)

    currency = show_col >= 0 and any(is_val(r) for r in aoa)

    tx: List[Dict[str, Any]] = []
    open_bal = close_bal = None
    n = len(aoa)
    for i in range(n):
        row = aoa[i] or []
        label = row[0] if (row and isinstance(row[0], str)) else ""
        nxt = aoa[i + 1] if i + 1 < n else None
        amt_row = (nxt or []) if (currency and is_val(nxt)) else row

        if re.search(r"Сальдо на начало", label, re.I):
            nums = [x for x in (_parse_num(c) for c in amt_row) if x is not None]
            if nums:
                open_bal = _signed_balance(amt_row, nums[-1])
            continue
        if re.search(r"Обороты за период|Сальдо на конец", label, re.I):
            nums = [x for x in (_parse_num(c) for c in amt_row) if x is not None]
            if nums:
                close_bal = _signed_balance(amt_row, nums[-1])
            continue

        date = _parse_date(row[0]) if row else None
        if not date:
            continue
        if currency and not is_val(nxt):
            continue  # валютная карточка: строка без «Вал.» (напр. «Переоценка»)

        deb = _parse_num(amt_row[deb_col]) if deb_col < len(amt_row) else None
        kre = _parse_num(amt_row[kre_col]) if kre_col < len(amt_row) else None
        doc = row[1] if len(row) > 1 else None
        no = _doc_no_1c(doc)
        if deb is not None and deb > 0:
            tx.append({"date": date, "dir": "in", "amount": _r2(deb), "no": no,
                       "party": _nth_line(row[4] if len(row) > 4 else None, 1),
                       "purpose": _nth_line(doc, 1), "parts": [_r2(deb)]})
        if kre is not None and kre > 0:
            tx.append({"date": date, "dir": "out", "amount": _r2(kre), "no": no,
                       "party": _nth_line(row[3] if len(row) > 3 else None, 1),
                       "purpose": _nth_line(doc, 1), "parts": [_r2(kre)]})
    return {"tx": tx, "open": open_bal, "close": close_bal, "currency": currency}


def _group_1c(lst: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    # Схлопываем разбитые проводки одного ПП (та же дата + направление + № док).
    m: Dict[str, Dict[str, Any]] = {}
    out: List[Dict[str, Any]] = []
    for t in lst:
        key = f'{t["date"]}|{t["dir"]}|{t["no"]}' if t["no"] else f"__{len(out)}"
        g = m.get(key)
        if g:
            g["amount"] = _r2(g["amount"] + t["amount"])
            g["parts"].append(t["amount"])
        else:
            copy = dict(t)
            copy["parts"] = list(t["parts"])
            m[key] = copy
            out.append(copy)
    return out


def _parse_bank(aoa: List[List[Cell]]) -> Dict[str, Any]:
    hdr_row, cols = _find_bank_cols(aoa)
    date_col = cols.get("date", 1)
    deb_col = cols.get("debit", 7)   # списание (out)
    kre_col = cols.get("credit", 8)  # приход (in)
    cp_col = cols.get("cp", 4)
    pur_col = cols.get("pur", 9)

    tx: List[Dict[str, Any]] = []
    open_bal = close_bal = None
    gaps: List[Dict[str, Any]] = []
    prev_close: Optional[float] = None
    prev_date = ""

    for row in aoa:
        row = row or []
        label = next((c for c in row if isinstance(c, str) and c.strip()), None)

        if label:
            ln = _norm(label)
            # Входящий остаток на ДАТА  /  Входящее сальдо  /  Кіріс сальдо
            m_in = re.search(r"входящий остаток на (\d{2}\.\d{2}\.\d{4})", ln)
            if m_in or "входящее сальдо" in ln or "кіріс сальдо" in ln:
                nums = [x for x in (_parse_num(c) for c in row) if x is not None]
                bal = nums[-1] if nums else None
                if open_bal is None:
                    open_bal = bal
                if m_in and bal is not None and prev_close is not None and _r2(bal) != _r2(prev_close):
                    gaps.append({"from_date": prev_date, "to_date": m_in.group(1),
                                 "amount": _r2(bal - prev_close)})
                continue
            # Исходящий остаток на ДАТА  /  Исходящее сальдо  /  Шығыс сальдо
            m_out = re.search(r"исходящий остаток на (\d{2}\.\d{2}\.\d{4})", ln)
            if m_out or "исходящее сальдо" in ln or "шығыс сальдо" in ln:
                nums = [x for x in (_parse_num(c) for c in row) if x is not None]
                bal = nums[-1] if nums else None
                if bal is not None:
                    close_bal = bal
                    prev_close = bal
                    prev_date = m_out.group(1) if m_out else prev_date
                continue
            # прочие служебные строки (итоги/шапки) без даты — пропускаем
            if (re.search(r"итого|жиынтығы|оборот|айналым|ведомость|№ док|№ п/п|реттік|входящ|исходящ|сальдо", ln)
                    and not _parse_date(row[date_col] if date_col < len(row) else None)):
                continue

        c0 = row[0] if row else None
        if c0 is None or not str(c0).strip():
            continue
        date = _parse_date(row[date_col] if date_col < len(row) else None)
        if not date:
            continue
        deb = _parse_num(row[deb_col]) if deb_col < len(row) else None
        kre = _parse_num(row[kre_col]) if kre_col < len(row) else None
        party = _first_line(row[cp_col] if cp_col < len(row) else None)
        pur_raw = row[pur_col] if pur_col < len(row) else None
        purpose = re.sub(r"\s+", " ", pur_raw).strip() if isinstance(pur_raw, str) else ""
        if deb is not None and deb > 0:
            tx.append({"date": date, "dir": "out", "amount": _r2(deb),
                       "no": str(c0).strip(), "party": party, "purpose": purpose, "parts": []})
        if kre is not None and kre > 0:
            tx.append({"date": date, "dir": "in", "amount": _r2(kre),
                       "no": str(c0).strip(), "party": party, "purpose": purpose, "parts": []})
    return {"tx": tx, "open": open_bal, "close": close_bal, "gaps": gaps}


class CaseBankService:
    """Сверка карточки счёта 1С ↔ банковская выписка."""

    def __init__(self, card_content: bytes, bank_content: bytes):
        self.card_content = card_content
        self.bank_content = bank_content

    def reconcile(self) -> Dict[str, Any]:
        c1 = _parse_1c(read_aoa(self.card_content))
        bank = _parse_bank(read_aoa(self.bank_content))
        c1g = _group_1c(c1["tx"])
        bank_tx = bank["tx"]

        rows: List[Dict[str, Any]] = []
        b_used = [False] * len(bank_tx)
        c_used = [False] * len(c1g)

        def match(same_date: bool):
            for a, b in enumerate(bank_tx):
                if b_used[a]:
                    continue
                for c, t in enumerate(c1g):
                    if c_used[c]:
                        continue
                    if b["dir"] != t["dir"]:
                        continue
                    if abs(b["amount"] - t["amount"]) > 0.005:
                        continue
                    if same_date and b["date"] != t["date"]:
                        continue
                    b_used[a] = True
                    c_used[c] = True
                    rows.append({
                        "date": b["date"], "date_c1": t["date"], "dir": b["dir"],
                        "amount": b["amount"], "bank_no": b["no"], "bank_party": b["party"],
                        "bank_purpose": b["purpose"], "c1_no": t["no"], "c1_party": t["party"],
                        "c1_purpose": t["purpose"], "parts": t["parts"],
                        "status": "ok" if same_date else "date_diff",
                    })
                    break

        match(True)
        match(False)

        for a, b in enumerate(bank_tx):
            if b_used[a]:
                continue
            rows.append({"date": b["date"], "dir": b["dir"], "amount": b["amount"],
                         "bank_no": b["no"], "bank_party": b["party"], "bank_purpose": b["purpose"],
                         "c1_no": "", "c1_party": "", "c1_purpose": "", "parts": [],
                         "status": "only_bank"})
        for c, t in enumerate(c1g):
            if c_used[c]:
                continue
            rows.append({"date": t["date"], "dir": t["dir"], "amount": t["amount"],
                         "bank_no": "", "bank_party": "", "bank_purpose": "",
                         "c1_no": t["no"], "c1_party": t["party"], "c1_purpose": t["purpose"],
                         "parts": t["parts"], "status": "only_1c"})

        rows.sort(key=lambda r: (_date_key(r["date"]), 0 if r["dir"] == "in" else 1))

        def s(arr, d):
            return _r2(sum(x["amount"] for x in arr if x["dir"] == d))

        bank_in, bank_out = s(bank_tx, "in"), s(bank_tx, "out")
        c1_in, c1_out = s(c1["tx"], "in"), s(c1["tx"], "out")

        open_c1, open_bank = c1["open"], bank["open"]
        close_bank = bank["close"]
        if c1["close"] is not None:
            close_c1 = c1["close"]
        elif open_c1 is not None:
            close_c1 = _r2(open_c1 + c1_in - c1_out)
        else:
            close_c1 = None
        balance_diff = (_r2(close_bank - close_c1)
                        if close_bank is not None and close_c1 is not None else None)

        split_docs = sorted({t["no"] for t in c1g if len(t["parts"]) > 1 and t["no"]})

        return {
            "rows": rows,
            "matched": sum(1 for r in rows if r["status"] in ("ok", "date_diff")),
            "only_bank": sum(1 for r in rows if r["status"] == "only_bank"),
            "only_1c": sum(1 for r in rows if r["status"] == "only_1c"),
            "bank_in": bank_in, "bank_out": bank_out, "c1_in": c1_in, "c1_out": c1_out,
            "open_bank": open_bank, "open_c1": open_c1,
            "close_bank": close_bank, "close_c1": close_c1,
            "balance_diff": balance_diff, "gaps": bank["gaps"],
            "split_docs": split_docs, "currency": c1["currency"],
        }


def _status_label(s: str) -> str:
    return {"ok": "OK", "date_diff": "Разная дата",
            "only_bank": "Нет в 1С", "only_1c": "Нет в банке"}.get(s, s)


def _fmt(n: Optional[float]) -> str:
    if n is None:
        return ""
    return f"{n:,.2f}".replace(",", " ")


def export_bank(res: Dict[str, Any]) -> bytes:
    wb = Workbook()

    # ---- Сводка ----
    ws1 = wb.active
    ws1.title = "Сводка"
    diff = res.get("balance_diff")
    S: List[List[Any]] = [["СВЕРКА БАНКОВСКОЙ ВЫПИСКИ С 1С"]]
    if res.get("currency"):
        S.append(["Валютный счёт: сверка по сумме в валюте (строка «Вал.»), тенговый эквивалент не участвует."])
    S += [
        [],
        ["Контрольная сверка", "Банк (выписка)", "1С (карточка счёта)", "Разница"],
        ["Входящий остаток", res.get("open_bank"), res.get("open_c1"),
         _r2(res["open_bank"] - res["open_c1"]) if res.get("open_bank") is not None and res.get("open_c1") is not None else ""],
        ["Обороты: приход", res.get("bank_in"), res.get("c1_in"), _r2(res.get("bank_in", 0) - res.get("c1_in", 0))],
        ["Обороты: списание", res.get("bank_out"), res.get("c1_out"), _r2(res.get("bank_out", 0) - res.get("c1_out", 0))],
        ["Исходящий остаток", res.get("close_bank"), res.get("close_c1"), diff if diff is not None else ""],
        ["Операций сопоставлено", res.get("matched"), "", ""],
        ["Нет в 1С (не проведено)", res.get("only_bank"), "", ""],
        ["Нет в банке", res.get("only_1c"), "", ""],
        [],
    ]
    if diff is not None and abs(diff) > 0.005:
        S.append(["⚠ РАСХОЖДЕНИЕ ПО ОСТАТКУ", _fmt(diff) + " ₸"])
        if res.get("gaps"):
            for g in res["gaps"]:
                S.append(["Причина",
                          f'Остаток в выписке изменился с {g["from_date"]} по {g["to_date"]} на {_fmt(g["amount"])} ₸ '
                          f'без операции (нет строки платежа). Эта сумма не отражена ни в выписке строкой, ни в 1С.'])
            S.append(["Что делать",
                      "Найти первичный документ на указанную сумму за этот период и провести его в 1С — остаток сравняется с банком."])
        else:
            S.append(["Причина",
                      "Обороты и/или остатки банка и 1С расходятся. Проверьте операции со статусом «Нет в 1С» и «Нет в банке» на листе «Детально»."])
    else:
        S.append(["РЕЗУЛЬТАТ", "Расхождений не найдено — обороты и остатки совпадают."])
    if res.get("split_docs"):
        S.append([])
        S.append(["Примечание",
                  "В 1С разбиты на 2 строки проводок (часть «Оплата», часть «Оплата (аванс)») платёжные поручения: №"
                  + ", №".join(res["split_docs"]) + ". Это не ошибка — сверка выполнена на уровне документа."])
    for r in S:
        ws1.append(r)
    ws1["A1"].font = Font(bold=True, size=13)
    for col, width in zip("ABCD", [30, 24, 26, 16]):
        ws1.column_dimensions[col].width = width

    # ---- Расхождения ----
    ws2 = wb.create_sheet("Расхождения")
    ws2.append(["№", "Дата", "Направление", "Сумма, ₸", "Тип расхождения", "Банк № / 1С №", "Комментарий"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    n = 1
    for g in res.get("gaps", []):
        ws2.append([n, f'{g["from_date"]}–{g["to_date"]}', "приход", g["amount"],
                    "Остаток вырос без операции", "", "Неучтённое поступление: провести в 1С"])
        n += 1
    for r in [x for x in res.get("rows", []) if x["status"] in ("only_bank", "only_1c", "date_diff")]:
        ws2.append([n, r["date"], "приход" if r["dir"] == "in" else "списание", r["amount"],
                    _status_label(r["status"]),
                    f'банк №{r["bank_no"]}' if r["status"] == "only_bank" else f'1С №{r["c1_no"]}',
                    (r.get("bank_purpose") or r.get("c1_purpose") or "")[:80]])
        n += 1
    if n == 1:
        ws2.append(["", "", "", "", "Расхождений не найдено", "", ""])
    for col, width in zip("ABCDEFG", [5, 18, 12, 16, 26, 16, 50]):
        ws2.column_dimensions[col].width = width

    # ---- Детально ----
    ws3 = wb.create_sheet("Детально")
    ws3.append(["Дата", "Направление", "Сумма, ₸", "Банк: № док", "Банк: контрагент",
                "Банк: назначение", "1С: № ПП", "1С: контрагент/статья", "1С: строк", "Статус"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for r in res.get("rows", []):
        parts = r.get("parts", [])
        ws3.append([
            r["date"], "приход" if r["dir"] == "in" else "списание", r["amount"],
            r.get("bank_no", ""), r.get("bank_party", ""), (r.get("bank_purpose") or "")[:80],
            r.get("c1_no", ""), r.get("c1_party") or r.get("c1_purpose", ""),
            f'2 ({" + ".join(_fmt(x) for x in parts)})' if len(parts) > 1 else "1",
            _status_label(r["status"]),
        ])
    for col, width in zip("ABCDEFGHIJ", [12, 11, 16, 11, 26, 44, 9, 24, 20, 16]):
        ws3.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
