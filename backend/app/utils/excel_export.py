import pandas as pd
from io import BytesIO
from typing import Dict, List, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows


class ExcelExporter:
    """Export reconciliation results to Excel with formatting"""

    COLORS = {
        'header': 'CCE5FF',
        'matched': 'C6EFCE',
        'mismatched': 'FFC7CE',
        'not_found': 'FFEB9C',
    }

    def __init__(self, results: Dict[str, Any], case_type: str):
        self.results = results
        self.case_type = case_type

    def export(self) -> bytes:
        """Export results to Excel file"""
        wb = Workbook()

        if self.case_type == 'case2' and 'header_summary' in self.results:
            # Case 2 (act reconciliation): Сводка / Сверка (общий период) / Вне периода
            ws_sum = wb.active
            ws_sum.title = "Сводка"
            self._add_act_overview(ws_sum)
            self._adjust_columns(ws_sum)

            ws_rec = wb.create_sheet("Сверка (общий период)")
            self._add_act_rows(ws_rec, self.results.get('data', []))
            self._adjust_columns(ws_rec)

            out_rows = self.results.get('out_of_period') or []
            if out_rows:
                ws_out = wb.create_sheet("Вне периода")
                self._add_act_rows(ws_out, out_rows)
                self._adjust_columns(ws_out)
        elif self.case_type == 'case2' and 'day_summaries' in self.results:
            # Legacy Case 2 result shape (day summaries) — kept for back-compat
            ws_days = wb.active
            ws_days.title = "Итоги по дням"
            self._add_day_summaries(ws_days)
            self._adjust_columns(ws_days)

            ws_trans = wb.create_sheet("Транзакции")
            self._add_summary(ws_trans)
            self._add_data(ws_trans)
            self._adjust_columns(ws_trans)
        else:
            # Case 1 or fallback: single sheet
            ws = wb.active
            ws.title = "Результаты сверки"
            self._add_summary(ws)
            self._add_data(ws)
            self._adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def _add_day_summaries(self, ws):
        """Add day summaries sheet for Case 2"""
        day_summaries = self.results.get('day_summaries', [])

        # Title
        ws['A1'] = "Итоги по дням"
        ws['A1'].font = Font(bold=True, size=14)

        # Headers
        headers = ['Дата', 'Наш дебет', 'Наш кредит', 'Дебет контрагента', 'Кредит контрагента', 'Разница дебет', 'Разница кредит', 'Статус']
        header_fill = PatternFill(start_color=self.COLORS['header'], fill_type="solid")
        header_font = Font(bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        start_row = 3
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        # Data rows
        for row_idx, day in enumerate(day_summaries, start_row + 1):
            status = "✓ Совпадает" if day['matches'] else "✗ Расхождение"
            row_data = [
                day['date'],
                day['our_debit'],
                day['our_credit'],
                day['cp_debit'],
                day['cp_credit'],
                day['debit_diff'],
                day['credit_diff'],
                status
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border

                # Color based on match status
                if day['matches']:
                    cell.fill = PatternFill(start_color=self.COLORS['matched'], fill_type="solid")
                else:
                    cell.fill = PatternFill(start_color=self.COLORS['mismatched'], fill_type="solid")

        # Add totals row
        if day_summaries:
            total_row = len(day_summaries) + start_row + 1
            ws.cell(row=total_row, column=1, value="ИТОГО").font = Font(bold=True)
            ws.cell(row=total_row, column=2, value=sum(d['our_debit'] for d in day_summaries)).font = Font(bold=True)
            ws.cell(row=total_row, column=3, value=sum(d['our_credit'] for d in day_summaries)).font = Font(bold=True)
            ws.cell(row=total_row, column=4, value=sum(d['cp_debit'] for d in day_summaries)).font = Font(bold=True)
            ws.cell(row=total_row, column=5, value=sum(d['cp_credit'] for d in day_summaries)).font = Font(bold=True)
            ws.cell(row=total_row, column=6, value=sum(d['debit_diff'] for d in day_summaries)).font = Font(bold=True)
            ws.cell(row=total_row, column=7, value=sum(d['credit_diff'] for d in day_summaries)).font = Font(bold=True)

            for col_idx in range(1, 8):
                ws.cell(row=total_row, column=col_idx).border = border

    # ------------------------- Case 2: act reconciliation -------------------------
    @staticmethod
    def _num(v):
        return v if isinstance(v, (int, float)) else None

    def _add_act_overview(self, ws):
        """Сводка sheet: owners, period, saldo/turnovers + differences, counts."""
        hs = self.results.get('header_summary', {}) or {}
        summary = self.results.get('summary', {}) or {}
        bold = Font(bold=True)
        title_font = Font(bold=True, size=14)
        header_fill = PatternFill(start_color=self.COLORS['header'], fill_type="solid")

        ws['A1'] = "СВЕРКА АКТОВ СВЕРКИ ВЗАИМОРАСЧЁТОВ"
        ws['A1'].font = title_font

        rows = [
            [],
            ["Акт 1", f"{hs.get('owner_act1') or '—'}", f"период {hs.get('period_act1') or '?'}"],
            ["Акт 2", f"{hs.get('owner_act2') or '—'}", f"период {hs.get('period_act2') or '?'}"],
            ["Правило", "Реализация ТМЗ №N (продавец) = Накладная/Товарный чек № вх. N (покупатель); сверка по номеру, ЭСФ, дате и сумме"],
        ]
        if hs.get('period_mismatch'):
            rows.append(["Период актов не совпадает", hs.get('period_act1'), hs.get('period_act2'),
                         f"Сопоставимы только за общий период {hs.get('common_period') or '—'}."])
        for r in rows:
            ws.append(r)

        ws.append([])
        head = ["Показатель", "Акт 1", "Акт 2", "Разница / комментарий"]
        ws.append(head)
        for c in range(1, 5):
            cell = ws.cell(row=ws.max_row, column=c)
            cell.font = bold
            cell.fill = header_fill

        odiff = self._num(hs.get('opening_diff'))
        cdiff = self._num(hs.get('closing_diff'))
        pm = hs.get('period_mismatch')
        data_rows = [
            ["Сальдо на начало", hs.get('opening_act1'), hs.get('opening_act2'),
             (f"разница {odiff:,.2f} ₸ (возникла до периода)".replace(',', ' ')
              if odiff not in (None, 0) else "совпадает")],
            ["Обороты Дебет", hs.get('turnover_debit_act1'), hs.get('turnover_debit_act2'),
             "разные периоды — не сравнивать" if pm else ""],
            ["Обороты Кредит", hs.get('turnover_credit_act1'), hs.get('turnover_credit_act2'),
             "разные периоды — не сравнивать" if pm else ""],
            ["Сальдо на конец", hs.get('closing_act1'), hs.get('closing_act2'),
             (f"разница {cdiff:,.2f} ₸".replace(',', ' ') if cdiff not in (None, 0) else "совпадает")],
        ]
        for r in data_rows:
            ws.append(r)

        ws.append([])
        counts = [
            ["Операций сверено (общий период)", summary.get('matched_count', self.results.get('matched', 0))],
            ["Расхождений (общий период)", self.results.get('mismatched', 0)],
            ["Нет у контрагента", self.results.get('not_found_in_source1', 0)],
            ["Нет у нас", self.results.get('not_found_in_source2', 0)],
            ["Строк вне общего периода", len(self.results.get('out_of_period') or [])],
            ["Совпадение", f"{summary.get('match_rate', 0)}%"],
            ["Автоопределение структуры", "да" if summary.get('auto_detected') else "ручная настройка"],
        ]
        for label, val in counts:
            ws.append([label, val])
            ws.cell(row=ws.max_row, column=1).font = bold

    def _add_act_rows(self, ws, data):
        """Сверка / Вне периода sheet: one row per operation, colored by status."""
        headers = ["Категория", "№ (акт 1)", "№ (акт 2)", "Дата", "Сумма акт 1", "Сумма акт 2",
                   "ЭСФ акт 1", "ЭСФ акт 2", "Метод", "Статус", "Комментарий", "Стр. 1", "Стр. 2"]
        keys = ["category", "num1", "num2", "date", "amount1", "amount2", "esf1", "esf2",
                "match_method", "status_detail", "note", "row1", "row2"]
        header_fill = PatternFill(start_color=self.COLORS['header'], fill_type="solid")
        header_font = Font(bold=True)
        border = Border(left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin'))

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        for row_idx, record in enumerate(data, 2):
            for col_idx, key in enumerate(keys, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=record.get(key))
                cell.border = border
            status = record.get('status', '')
            fill = None
            if status == 'Совпадает':
                fill = PatternFill(start_color=self.COLORS['matched'], fill_type="solid")
            elif status == 'Расхождение':
                fill = PatternFill(start_color=self.COLORS['mismatched'], fill_type="solid")
            elif status in ('Нет у контрагента', 'Нет у нас', 'Вне периода'):
                fill = PatternFill(start_color=self.COLORS['not_found'], fill_type="solid")
            if fill:
                for col_idx in range(1, len(headers) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = fill

    def _add_summary(self, ws):
        """Add summary section at the top"""
        summary_data = [
            ["Результаты сверки"],
            [""],
            ["Всего записей:", self.results['total_records']],
            ["Совпадает:", self.results['matched']],
            ["Расхождения:", self.results['mismatched']],
            ["Не найдено в источнике 1:", self.results['not_found_in_source1']],
            ["Не найдено в источнике 2:", self.results['not_found_in_source2']],
            [""],
        ]

        header_font = Font(bold=True, size=14)
        ws['A1'].font = header_font

        for row_idx, row_data in enumerate(summary_data, 1):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

    def _add_data(self, ws):
        """Add reconciliation data"""
        data = self.results['data']
        if not data:
            return

        start_row = 9
        df = pd.DataFrame(data)

        # Column headers
        if self.case_type == 'case1':
            columns = {
                'date': 'Дата',
                'document': 'Документ',
                'counterparty_6010': 'Контрагент (6010)',
                'recipient_esf': 'Получатель (ЭСФ)',
                'amount_6010': 'Сумма 6010',
                'amount_esf': 'Сумма ЭСФ',
                'esf_number': '№ ЭСФ',
                'esf_reg_number': 'Рег. номер ЭСФ',
                'difference': 'Разница',
                'status': 'Статус'
            }
        elif self.case_type == 'case3':
            columns = {
                'date': 'Дата',
                'document': 'Документ',
                'tru': 'ТРУ',
                'sender_esf': 'Отправитель (ЭСФ)',
                'recipient_esf': 'Получатель (ЭСФ)',
                'amount_3310': 'Сумма 3310',
                'amount_esf': 'Сумма ЭСФ',
                'esf_number': '№ ЭСФ',
                'esf_reg_number': 'Рег. номер ЭСФ',
                'difference': 'Разница',
                'status': 'Статус'
            }
        else:  # case2
            columns = {
                'date': 'Дата',
                'document': 'Наш документ',
                'our_doc_number': '№ нашего док.',
                'cp_document': 'Документ контрагента',
                'cp_doc_number': '№ док. контрагента',
                'our_debit': 'Наш дебет',
                'our_credit': 'Наш кредит',
                'cp_debit': 'Дебет контрагента',
                'cp_credit': 'Кредит контрагента',
                'debit_diff': 'Разница дебет',
                'credit_diff': 'Разница кредит',
                'match_phase': 'Метод сопоставления',
                'status': 'Статус'
            }

        # Write headers
        header_fill = PatternFill(start_color=self.COLORS['header'], fill_type="solid")
        header_font = Font(bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col_idx, (key, header) in enumerate(columns.items(), 1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        # Write data rows
        for row_idx, record in enumerate(data, start_row + 1):
            for col_idx, key in enumerate(columns.keys(), 1):
                value = record.get(key)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border

                # Color based on status
                status = record.get('status', '')
                if status == 'Совпадает':
                    cell.fill = PatternFill(start_color=self.COLORS['matched'], fill_type="solid")
                elif status == 'Расхождение':
                    cell.fill = PatternFill(start_color=self.COLORS['mismatched'], fill_type="solid")
                elif 'Не найдено' in status:
                    cell.fill = PatternFill(start_color=self.COLORS['not_found'], fill_type="solid")

    def _adjust_columns(self, ws):
        """Auto-adjust column widths"""
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
